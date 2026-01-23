"""Gerador de documentos para laudos técnicos."""

import logging
from pathlib import Path
from typing import Dict, Any, Optional
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from .report_templates import ReportTemplates

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class DocumentGenerator:
    """Gerador de documentos em múltiplos formatos."""

    def __init__(self, output_dir: Optional[Path] = None):
        """
        Inicializa o gerador.

        Args:
            output_dir: Diretório de saída para documentos
        """
        self.output_dir = output_dir or Path("data/output")
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.templates = ReportTemplates()

    def generate_full_report(
        self,
        analysis_data: Dict[str, Any],
        format: str = "txt",
        filename: Optional[str] = None
    ) -> Path:
        """
        Gera laudo técnico completo.

        Args:
            analysis_data: Dados da análise
            format: Formato do documento (txt, docx, pdf)
            filename: Nome do arquivo (opcional)

        Returns:
            Path do arquivo gerado
        """
        try:
            if not filename:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"laudo_tecnico_{timestamp}"

            template = self.templates.get_full_report_template()

            # Preencher template
            report_text = template.format(**self._prepare_report_data(analysis_data))

            if format == "txt":
                return self._save_as_txt(report_text, filename)
            elif format == "docx":
                return self._save_as_docx(report_text, filename, analysis_data)
            elif format == "pdf":
                return self._save_as_pdf(report_text, filename)
            else:
                raise ValueError(f"Formato não suportado: {format}")

        except Exception as e:
            logger.error(f"Erro ao gerar relatório: {str(e)}")
            raise

    def generate_summary(
        self,
        analysis_data: Dict[str, Any],
        format: str = "txt",
        filename: Optional[str] = None
    ) -> Path:
        """
        Gera resumo executivo.

        Args:
            analysis_data: Dados da análise
            format: Formato do documento
            filename: Nome do arquivo

        Returns:
            Path do arquivo gerado
        """
        try:
            if not filename:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"resumo_executivo_{timestamp}"

            template = self.templates.get_summary_template()
            summary_data = self._prepare_summary_data(analysis_data)
            summary_text = template.format(**summary_data)

            if format == "txt":
                return self._save_as_txt(summary_text, filename)
            elif format == "docx":
                return self._save_as_docx(summary_text, filename, analysis_data)
            elif format == "pdf":
                return self._save_as_pdf(summary_text, filename)

        except Exception as e:
            logger.error(f"Erro ao gerar resumo: {str(e)}")
            raise

    def generate_risk_assessment(
        self,
        analysis_data: Dict[str, Any],
        format: str = "txt",
        filename: Optional[str] = None
    ) -> Path:
        """
        Gera avaliação de riscos.

        Args:
            analysis_data: Dados da análise
            format: Formato do documento
            filename: Nome do arquivo

        Returns:
            Path do arquivo gerado
        """
        try:
            if not filename:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"avaliacao_riscos_{timestamp}"

            template = self.templates.get_risk_assessment_template()
            risk_data = self._prepare_risk_data(analysis_data)
            risk_text = template.format(**risk_data)

            if format == "txt":
                return self._save_as_txt(risk_text, filename)
            elif format == "docx":
                return self._save_as_docx(risk_text, filename, analysis_data)
            elif format == "pdf":
                return self._save_as_pdf(risk_text, filename)

        except Exception as e:
            logger.error(f"Erro ao gerar avaliação de riscos: {str(e)}")
            raise

    def _prepare_report_data(self, analysis: Dict[str, Any]) -> Dict[str, str]:
        """Prepara dados para preenchimento do template."""
        now = datetime.now()

        return {
            "laudo_numero": analysis.get("laudo_numero", f"LAU{now.strftime('%Y%m%d%H%M%S')}"),
            "data_emissao": now.strftime("%d/%m/%Y"),
            "perito_nome": analysis.get("perito_nome", "A DEFINIR"),
            "perito_registro": analysis.get("perito_registro", "A DEFINIR"),
            "perito_especialidade": analysis.get("perito_especialidade", "Direito Contratual"),
            "parte1_nome": analysis.get("parte1", {}).get("nome", "A DEFINIR"),
            "parte1_documento": analysis.get("parte1", {}).get("documento", "A DEFINIR"),
            "parte1_endereco": analysis.get("parte1", {}).get("endereco", "A DEFINIR"),
            "parte1_representante": analysis.get("parte1", {}).get("representante", "A DEFINIR"),
            "parte2_nome": analysis.get("parte2", {}).get("nome", "A DEFINIR"),
            "parte2_documento": analysis.get("parte2", {}).get("documento", "A DEFINIR"),
            "parte2_endereco": analysis.get("parte2", {}).get("endereco", "A DEFINIR"),
            "parte2_representante": analysis.get("parte2", {}).get("representante", "A DEFINIR"),
            "objeto_pericia": analysis.get("objeto_pericia", "Análise de contrato jurídico"),
            "objetivo": analysis.get("objetivo", "Análise técnica do contrato apresentado"),
            "quesitos": analysis.get("quesitos", "A DEFINIR"),
            "metodologia": analysis.get("metodologia", "Análise documental com suporte de IA"),
            "documentos_analisados": analysis.get("documentos_analisados", "Contrato em PDF"),
            "tipo_contrato": analysis.get("tipo_contrato", "A DEFINIR"),
            "data_assinatura": analysis.get("data_assinatura", "A DEFINIR"),
            "local_assinatura": analysis.get("local_assinatura", "A DEFINIR"),
            "num_paginas": str(analysis.get("num_paginas", "N/A")),
            "num_clausulas": str(analysis.get("num_clausulas", "N/A")),
            "objeto_contratual": analysis.get("objeto_contratual", "A DEFINIR"),
            "analise_clausulas": analysis.get("analise_clausulas", "A DEFINIR"),
            "obrigacoes_parte1": analysis.get("obrigacoes_parte1", "A DEFINIR"),
            "obrigacoes_parte2": analysis.get("obrigacoes_parte2", "A DEFINIR"),
            "conformidade_legal": analysis.get("conformidade_legal", "A DEFINIR"),
            "base_legal": analysis.get("base_legal", "Código Civil Brasileiro, CDC"),
            "clausulas_atencao": analysis.get("clausulas_atencao", "A DEFINIR"),
            "riscos_juridicos": analysis.get("riscos_juridicos", "A DEFINIR"),
            "valor_total": analysis.get("valor_total", "A DEFINIR"),
            "forma_pagamento": analysis.get("forma_pagamento", "A DEFINIR"),
            "parcelas": analysis.get("parcelas", "A DEFINIR"),
            "reajustes": analysis.get("reajustes", "A DEFINIR"),
            "multas_penalidades": analysis.get("multas_penalidades", "A DEFINIR"),
            "garantias": analysis.get("garantias", "A DEFINIR"),
            "projecao_financeira": analysis.get("projecao_financeira", "A DEFINIR"),
            "prazo_vigencia": analysis.get("prazo_vigencia", "A DEFINIR"),
            "condicoes_prorrogacao": analysis.get("condicoes_prorrogacao", "A DEFINIR"),
            "condicoes_rescisao": analysis.get("condicoes_rescisao", "A DEFINIR"),
            "prazos_prescricionais": analysis.get("prazos_prescricionais", "A DEFINIR"),
            "matriz_riscos": analysis.get("matriz_riscos", "A DEFINIR"),
            "detalhamento_riscos_juridicos": analysis.get("detalhamento_riscos_juridicos", "A DEFINIR"),
            "detalhamento_riscos_financeiros": analysis.get("detalhamento_riscos_financeiros", "A DEFINIR"),
            "riscos_operacionais": analysis.get("riscos_operacionais", "A DEFINIR"),
            "respostas_quesitos": analysis.get("respostas_quesitos", "A DEFINIR"),
            "sintese_analise": analysis.get("sintese_analise", "A DEFINIR"),
            "parecer_tecnico": analysis.get("parecer_tecnico", "A DEFINIR"),
            "recomendacoes": analysis.get("recomendacoes", "A DEFINIR"),
            "local_data_final": f"{analysis.get('local', 'São Paulo')}, {now.strftime('%d de %B de %Y')}",
            "anexos": analysis.get("anexos", "Não há anexos")
        }

    def _prepare_summary_data(self, analysis: Dict[str, Any]) -> Dict[str, str]:
        """Prepara dados para resumo."""
        return {
            "data": datetime.now().strftime("%d/%m/%Y"),
            "tipo_contrato": analysis.get("tipo_contrato", "A DEFINIR"),
            "partes": f"{analysis.get('parte1', {}).get('nome', 'N/A')} e {analysis.get('parte2', {}).get('nome', 'N/A')}",
            "objeto": analysis.get("objeto_contratual", "A DEFINIR"),
            "valor_total": analysis.get("valor_total", "A DEFINIR"),
            "prazo": analysis.get("prazo_vigencia", "A DEFINIR"),
            "vigencia": analysis.get("prazo_vigencia", "A DEFINIR"),
            "obrigacoes": analysis.get("sintese_obrigacoes", "A DEFINIR"),
            "pontos_criticos": analysis.get("clausulas_atencao", "A DEFINIR"),
            "riscos": analysis.get("riscos_juridicos", "A DEFINIR"),
            "recomendacoes": analysis.get("recomendacoes", "A DEFINIR")
        }

    def _prepare_risk_data(self, analysis: Dict[str, Any]) -> Dict[str, str]:
        """Prepara dados para avaliação de riscos."""
        return {
            "contrato_id": analysis.get("laudo_numero", "N/A"),
            "data": datetime.now().strftime("%d/%m/%Y"),
            "matriz_riscos_tabela": analysis.get("matriz_riscos", "A DEFINIR"),
            "nivel_risco_juridico": analysis.get("nivel_risco_juridico", "MÉDIO"),
            "nivel_risco_financeiro": analysis.get("nivel_risco_financeiro", "MÉDIO"),
            "nivel_risco_operacional": analysis.get("nivel_risco_operacional", "BAIXO"),
            "riscos_juridicos": analysis.get("detalhamento_riscos_juridicos", "A DEFINIR"),
            "riscos_financeiros": analysis.get("detalhamento_riscos_financeiros", "A DEFINIR"),
            "riscos_operacionais": analysis.get("riscos_operacionais", "A DEFINIR"),
            "plano_mitigacao": analysis.get("plano_mitigacao", "A DEFINIR"),
            "nivel_risco_geral": analysis.get("nivel_risco_geral", "MÉDIO"),
            "conclusao_riscos": analysis.get("conclusao_riscos", "A DEFINIR")
        }

    def _save_as_txt(self, content: str, filename: str) -> Path:
        """Salva como arquivo de texto."""
        output_path = self.output_dir / f"{filename}.txt"
        output_path.write_text(content, encoding='utf-8')
        logger.info(f"Relatório TXT salvo em: {output_path}")
        return output_path

    def _save_as_docx(self, content: str, filename: str, data: Dict[str, Any]) -> Path:
        """Salva como documento Word."""
        output_path = self.output_dir / f"{filename}.docx"

        doc = Document()

        # Configurar estilos
        style = doc.styles['Normal']
        style.font.name = 'Arial'
        style.font.size = Pt(11)

        # Título
        title = doc.add_heading('LAUDO TÉCNICO PERICIAL', 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Adicionar conteúdo
        for line in content.split('\n'):
            if line.startswith('═'):
                continue  # Pular linhas decorativas
            elif line.strip().startswith(tuple(str(i) for i in range(10))):
                # Cabeçalhos de seção
                doc.add_heading(line.strip(), level=1)
            else:
                doc.add_paragraph(line)

        doc.save(str(output_path))
        logger.info(f"Relatório DOCX salvo em: {output_path}")
        return output_path

    def _save_as_pdf(self, content: str, filename: str) -> Path:
        """Salva como PDF."""
        output_path = self.output_dir / f"{filename}.pdf"

        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=18
        )

        styles = getSampleStyleSheet()
        story = []

        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=RGBColor(0, 0, 0),
            spaceAfter=30,
            alignment=1  # Center
        )

        # Adicionar conteúdo
        for line in content.split('\n'):
            if line.startswith('═'):
                continue
            elif line.strip():
                story.append(Paragraph(line, styles['Normal']))
                story.append(Spacer(1, 12))

        doc.build(story)
        logger.info(f"Relatório PDF salvo em: {output_path}")
        return output_path
